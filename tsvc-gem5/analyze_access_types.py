#!/usr/bin/env python3
"""Static RVV memory-access-type profile of TSVC kernels from an objdump dump.

Reads `objdump -d` output (e.g. bin/GNU/tsvc_vec.dump), and for every TSVC
kernel function counts the vector load/store instructions by addressing-mode
family. Because it reads the *emitted* instructions, the result is immune to
guessing wrong about what the compiler actually vectorized.

Output: an .xlsx with
  - "Counts": absolute static counts per kernel, every load/store category.
  - "Chart":  stacked column chart of the access-type mix per vectorized kernel.

NOTE: these are STATIC instruction-site counts (how many of each instruction
appear in the kernel body), not dynamic execution counts. For frequency-
weighted, executed counts use gem5's committedInstType::Simd* stats in
stats.txt (scoped per kernel by the ROI m5_reset_stats / m5_dump_reset_stats).

Usage:
    ./analyze_access_types.py bin/GNU/tsvc_vec.dump [-o tsvc_vec_access_types.xlsx]
"""
import argparse
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Classification: ordered rules, first match wins. Order matters because the
# mnemonics share prefixes (segment vs strided vs unit). Direction (load/store)
# comes from the 2nd letter: 'l' = load, 's' = store.
# ---------------------------------------------------------------------------
RULES = [
    ("seg-indexed", re.compile(r"^v[ls](ux|ox)seg")),   # vluxseg/vloxseg/vsuxseg/vsoxseg
    ("seg-strided", re.compile(r"^v[ls]sseg")),          # vlsseg / vssseg
    ("seg-unit",    re.compile(r"^v[ls]seg")),           # vlseg / vsseg (+ff)
    ("indexed",     re.compile(r"^v[ls](ux|ox)ei\d")),   # vluxei/vloxei/vsuxei/vsoxei
    ("strided",     re.compile(r"^v[ls]se\d")),          # vlse32 / vsse32
    ("whole-reg",   re.compile(r"^v[ls]\d+re?\d*\.")),   # vl1re32 / vs1r  (spill/fill)
    ("unit",        re.compile(r"^v[ls]e\d")),           # vle32 / vse32 / vle8ff
    ("unit",        re.compile(r"^v[ls]m\.")),           # vlm.v / vsm.v (mask)
]

# Column order for the output table. Each base category -> Load + Store column.
BASE_CATS = ["unit", "strided", "indexed",
             "seg-unit", "seg-strided", "seg-indexed", "whole-reg"]
PRETTY = {
    "unit": "Unit-Stride", "strided": "Strided (const)", "indexed": "Indexed",
    "seg-unit": "Seg Unit-Stride", "seg-strided": "Seg Strided",
    "seg-indexed": "Seg Indexed", "whole-reg": "Whole-Reg (spill)",
}
# Categories that represent real data access patterns (whole-reg is spill noise).
DATA_CATS = ["unit", "strided", "indexed", "seg-unit", "seg-strided", "seg-indexed"]

SYM_RE = re.compile(r"^[0-9a-f]+ <([^>]+)>:")
INSN_RE = re.compile(r"^\s+[0-9a-f]+:\t")


def classify(mnem):
    """Return (base_category, 'L'|'S') or None if not a vector load/store."""
    for cat, rx in RULES:
        if rx.match(mnem):
            return cat, ("L" if mnem[1] == "l" else "S")
    return None


def parse_functions(path):
    """dump -> {func_name: {'mnems': [..], 'vec_any': bool}}."""
    funcs, cur = {}, None
    with open(path) as f:
        for line in f:
            m = SYM_RE.match(line)
            if m:
                cur = m.group(1)
                funcs[cur] = {"mnems": [], "vec_any": False}
                continue
            if cur is None or not INSN_RE.match(line):
                continue
            parts = line.split("\t")
            if len(parts) < 3:
                continue
            mnem = parts[2].split()[0] if parts[2].strip() else ""
            if not mnem:
                continue
            funcs[cur]["mnems"].append(mnem)
            if re.match(r"^v[a-z]", mnem):
                funcs[cur]["vec_any"] = True
    return funcs


def load_kernel_list(repo_dir):
    """Extract the kernel name list from run-tsvc.sh (programs=( ... ))."""
    path = os.path.join(repo_dir, "run-tsvc.sh")
    text = open(path).read()
    m = re.search(r"programs=\((.*?)\)", text, re.S)
    if not m:
        return None
    return re.findall(r'"([^"]+)"', m.group(1))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dump", help="objdump -d output (e.g. bin/GNU/tsvc_vec.dump)")
    ap.add_argument("-o", "--out", default=None, help="output .xlsx path")
    args = ap.parse_args()

    repo_dir = os.path.dirname(os.path.abspath(args.dump)) or "."
    # walk up to find run-tsvc.sh (dump is in bin/GNU/)
    search = os.path.abspath(args.dump)
    kernels = None
    for _ in range(4):
        search = os.path.dirname(search)
        if os.path.exists(os.path.join(search, "run-tsvc.sh")):
            kernels = load_kernel_list(search)
            break
    if not kernels:
        print("ERROR: could not find run-tsvc.sh kernel list", file=sys.stderr)
        return 1

    funcs = parse_functions(args.dump)

    # Build per-kernel category counts.
    rows = []
    for k in kernels:
        if k not in funcs:
            continue  # kernel not compiled into this binary
        counts = {(c, d): 0 for c in BASE_CATS for d in ("L", "S")}
        for mnem in funcs[k]["mnems"]:
            r = classify(mnem)
            if r:
                counts[r] += 1
        total = sum(counts.values())
        data_total = sum(counts[(c, d)] for c in DATA_CATS for d in ("L", "S"))
        present = [PRETTY[c] for c in DATA_CATS
                   if counts[(c, "L")] or counts[(c, "S")]]
        if data_total == 0:
            cls = "scalar (not vectorized)" if not funcs[k]["vec_any"] \
                  else "vector compute, scalar mem"
        else:
            cls = " + ".join(present)
        rows.append((k, counts, total, cls))

    out = args.out or os.path.splitext(os.path.basename(args.dump))[0] + \
        "_access_types.xlsx"

    write_xlsx(out, rows)

    # stdout summary
    nvec = sum(1 for _, c, _, _ in rows
               if any(c[(cat, d)] for cat in DATA_CATS for d in ("L", "S")))
    print(f"Parsed {len(rows)} kernels from {args.dump}")
    print(f"  vectorized (>=1 vector data load/store): {nvec}")
    print(f"  scalar / not vectorized:                 {len(rows) - nvec}")
    cat_tot = {c: sum(cnt[(c, d)] for _, cnt, _, _ in rows for d in ("L", "S"))
               for c in BASE_CATS}
    print("  static instruction totals by category:")
    for c in BASE_CATS:
        print(f"    {PRETTY[c]:<20} {cat_tot[c]}")
    print(f"Wrote {out}")
    return 0


def write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Counts"

    # Header: two-row grouped header (category over Load/Store).
    hdr1 = ["Kernel"]
    hdr2 = [""]
    col_keys = []
    for c in BASE_CATS:
        hdr1 += [PRETTY[c], ""]
        hdr2 += ["Ld", "St"]
        col_keys += [(c, "L"), (c, "S")]
    hdr1 += ["Total", "Classification"]
    hdr2 += ["", ""]

    ws.append(hdr1)
    ws.append(hdr2)

    bold = Font(bold=True)
    center = Alignment(horizontal="center")
    fill = PatternFill("solid", fgColor="DDEBF7")
    for r in (1, 2):
        for cell in ws[r]:
            cell.font = bold
            cell.alignment = center
            cell.fill = fill
    # merge category headers across their Ld/St pair
    col = 2
    for c in BASE_CATS:
        ws.merge_cells(start_row=1, start_column=col,
                       end_row=1, end_column=col + 1)
        col += 2
    ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)      # Total
    ws.merge_cells(start_row=1, start_column=col + 1, end_row=2, end_column=col + 1)  # Class

    for k, counts, total, cls in rows:
        row = [k] + [counts[key] for key in col_keys] + [total, cls]
        ws.append(row)

    # widths
    ws.column_dimensions["A"].width = 10
    for i in range(2, 2 + len(col_keys)):
        ws.column_dimensions[get_column_letter(i)].width = 6
    tot_col = 2 + len(col_keys)
    ws.column_dimensions[get_column_letter(tot_col)].width = 7
    ws.column_dimensions[get_column_letter(tot_col + 1)].width = 28
    ws.freeze_panes = "B3"

    build_chart(wb, ws, rows, col_keys)
    wb.save(path)


def build_chart(wb, counts_ws, rows, col_keys):
    """Stacked column chart of access-type mix, vectorized kernels only,
    non-empty data categories only."""
    # Which data categories are non-empty across the dataset?
    used = []
    for idx, (c, d) in enumerate(col_keys):
        if c not in DATA_CATS:
            continue
        if any(cnt[(c, d)] for _, cnt, _, _ in rows):
            used.append((idx, c, d))

    vec_rows = [(k, cnt) for k, cnt, _, _ in rows
                if any(cnt[(c, d)] for c in DATA_CATS for d in ("L", "S"))]

    cs = wb.create_sheet("ChartData")
    cs.append(["Kernel"] + [f"{PRETTY[c]} {d}" for _, c, d in used])
    for k, cnt in vec_rows:
        cs.append([k] + [cnt[(c, d)] for _, c, d in used])

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "RVV memory-access-type mix per TSVC kernel (static counts)"
    chart.y_axis.title = "vector load/store instructions"
    chart.x_axis.title = "kernel"
    chart.height = 12
    chart.width = max(30, 0.55 * len(vec_rows))

    nrows = len(vec_rows)
    ncols = len(used)
    data = Reference(cs, min_col=2, max_col=1 + ncols, min_row=1, max_row=1 + nrows)
    cats = Reference(cs, min_col=1, min_row=2, max_row=1 + nrows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.gapWidth = 30

    chart_ws = wb.create_sheet("Chart")
    chart_ws.add_chart(chart, "A1")


if __name__ == "__main__":
    sys.exit(main())
